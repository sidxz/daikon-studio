"""End-to-end tests for the Run routes: submit a prediction, poll it, read its
results, cancel it.

Training and prediction both run through the real `InlineEnqueuer` (see the
`app` fixture in `tests/api/conftest.py`), so by the time `POST /api/v1/runs`
returns, the prediction it names has already finished -- the same "no polling
needed" shape `test_protocols.py` documents for training.
"""

from __future__ import annotations

import uuid
from typing import Any
from urllib.parse import quote

import pytest
import pytest_asyncio
from sqlalchemy import update
from tests.fakes.tunable_data import tunable_csv

from daikonstudio.application.engines.checkpoints import checkpoint_root
from daikonstudio.application.execution.train_protocol import TrainProtocolCommand
from daikonstudio.application.ports.blob_store import BlobStore
from daikonstudio.domain.execution.run import Run, RunKind, RunStatus
from daikonstudio.infrastructure.persistence.sqlalchemy.catalog.models import InSilicoProtocolModel
from daikonstudio.infrastructure.persistence.sqlalchemy.execution.repository import (
    SqlAlchemyRunRepository,
)

# Twenty distinct compounds to train on -- the same shape `test_protocols.py`
# uses and for the same reason: a real 16/2/2 random split.
_STRUCTURES = (
    "CCO",
    "CCN",
    "CCCO",
    "CCCCO",
    "CCCCCO",
    "c1ccccc1",
    "Cc1ccccc1",
    "c1ccncc1",
    "c1ccsc1",
    "c1cc[nH]c1",
    "C1CCCCC1",
    "C1CCNCC1",
    "C1CCOC1",
    "C1CCCC1",
    "C1CC1",
    "c1ccc2ccccc2c1",
    "c1ccc2[nH]ccc2c1",
    "C1CCC2CCCCC2C1",
    "c1cnc2ccccc2c1",
    "O=C1CCCCC1",
)

# A colleague's own compounds to run the Protocol against.
_PREDICTION_CSV = b"smiles\nCCO\nc1ccccc1\nFc1ccc(F)cc1\n"


def _training_csv() -> bytes:
    rows = "\n".join(f"{smiles},{1.0 + 0.37 * index}" for index, smiles in enumerate(_STRUCTURES))
    return f"smiles,y\n{rows}\n".encode()


async def _create_dataset(client, csv_upload) -> str:
    upload_ref = await csv_upload(_training_csv())
    response = await client.post(
        "/api/v1/datasets",
        json={
            "name": "solubility",
            "upload_ref": upload_ref,
            "structure_column": "smiles",
            "targets": [{"column": "y", "kind": "numeric", "unit": "logS", "direction": "high"}],
            "split": {"strategy": "random", "seed": 1},
        },
    )
    assert response.status_code == 201, response.text
    return str(response.json()["id"])


async def _train(client, dataset_id: str, **overrides: object) -> Any:
    body: dict[str, object] = {
        "name": "solubility model",
        "dataset_id": dataset_id,
        "engine_id": "ecfp4-xgboost",
        "conditions": {},
    }
    body.update(overrides)
    return await client.post("/api/v1/protocols", json=body)


def _predict_body(protocol_id: str, upload_ref: str, **overrides: object) -> dict[str, object]:
    body: dict[str, object] = {
        "protocol_id": protocol_id,
        "upload_ref": upload_ref,
        "structure_column": "smiles",
        "conditions": {},
    }
    body.update(overrides)
    return body


async def _predict(client, protocol_id: str, upload_ref: str, **overrides: object) -> Any:
    return await client.post(
        "/api/v1/runs", json=_predict_body(protocol_id, upload_ref, **overrides)
    )


@pytest_asyncio.fixture
async def trained_protocol_id(client, csv_upload) -> str:
    dataset_id = await _create_dataset(client, csv_upload)
    response = await _train(client, dataset_id)
    assert response.status_code == 202, response.text
    listing = await client.get("/api/v1/protocols")
    items = listing.json()["items"]
    assert len(items) == 1, items
    return str(items[0]["id"])


@pytest_asyncio.fixture
async def published_protocol_id(client, trained_protocol_id) -> str:
    response = await client.post(f"/api/v1/protocols/{trained_protocol_id}/publish")
    assert response.status_code == 204, response.text
    return trained_protocol_id


@pytest_asyncio.fixture
async def prediction_upload_ref(csv_upload) -> str:
    return await csv_upload(_PREDICTION_CSV)


async def test_prediction_request_returns_202_and_is_immediately_ready(
    client, published_protocol_id, prediction_upload_ref
):
    response = await _predict(client, published_protocol_id, prediction_upload_ref)
    assert response.status_code == 202, response.text
    body = response.json()
    assert body["kind"] == "prediction"

    polled = await client.get(f"/api/v1/runs/{body['id']}")
    assert polled.status_code == 200, polled.text
    assert polled.json()["status"] == "ready"


async def test_identical_prediction_requests_reuse_the_cached_run(
    client, published_protocol_id, prediction_upload_ref
):
    first = await _predict(client, published_protocol_id, prediction_upload_ref)
    second = await _predict(client, published_protocol_id, prediction_upload_ref)
    assert first.status_code == 202, first.text
    assert second.status_code == 202, second.text
    assert first.json()["id"] == second.json()["id"]


async def test_results_carry_structure_readouts_uncertainty_and_applicability(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    run_id = submitted.json()["id"]

    protocol = (await client.get(f"/api/v1/protocols/{published_protocol_id}")).json()
    readout = protocol["readouts"][0]

    response = await client.get(f"/api/v1/runs/{run_id}/results")
    assert response.status_code == 200, response.text
    body = response.json()
    assert len(body["items"]) == 3  # all three query compounds canonicalize
    row = body["items"][0]
    assert "structure" in row
    assert readout["name"] in row["readouts"]
    # Important 5 (Task 17 review): unit and direction travel with the value,
    # not just a bare number -- otherwise a client can't line a prediction up
    # against a measurement without a second call to fetch the Protocol.
    predicted = row["readouts"][readout["name"]]
    assert isinstance(predicted["value"], float)
    assert predicted["unit"] == readout["unit"]
    assert predicted["direction"] == readout["direction"]
    assert "uncertainty" in row
    # ecfp4-xgboost: no ensemble spread to report, never a fabricated number.
    assert row["uncertainty"] == {"y": None}
    assert isinstance(row["applicability"], float)
    assert 0.0 <= row["applicability"] <= 1.0


async def test_a_one_target_results_file_keeps_the_plain_uncertainty_column(
    client, published_protocol_id, prediction_upload_ref
):
    """Every results file written before several targets existed has `uncertainty`,
    not `y_uncertainty`; a one-target protocol must keep writing and sorting by it."""
    run_id = (await _predict(client, published_protocol_id, prediction_upload_ref)).json()["id"]
    response = await client.get(
        f"/api/v1/runs/{run_id}/results", params={"sort_by": "uncertainty"}
    )
    assert response.status_code == 200, response.text


async def test_result_ranges_span_the_whole_run_for_every_numeric_column(
    client, published_protocol_id, prediction_upload_ref
):
    """The scale a triage grid draws its bars against: every readout, the uncertainty
    and applicability, over all rows, whatever page or filter the grid is showing."""
    run_id = (await _predict(client, published_protocol_id, prediction_upload_ref)).json()["id"]
    readout = (await client.get(f"/api/v1/protocols/{published_protocol_id}")).json()["readouts"][
        0
    ]["name"]
    rows = (await client.get(f"/api/v1/runs/{run_id}/results")).json()["items"]

    response = await client.get(f"/api/v1/runs/{run_id}/results/ranges")
    assert response.status_code == 200, response.text
    ranges = response.json()

    values = [row["readouts"][readout]["value"] for row in rows]
    assert ranges[readout] == {
        "min": pytest.approx(min(values)),
        "max": pytest.approx(max(values)),
    }
    similarity = [row["applicability"] for row in rows]
    assert ranges["applicability"] == {
        "min": pytest.approx(min(similarity)),
        "max": pytest.approx(max(similarity)),
    }
    # ecfp4-xgboost reports no uncertainty: a column with no values has no range.
    assert ranges["uncertainty"] == {"min": None, "max": None}


async def test_result_ranges_for_a_training_run_is_a_404(client, csv_upload):
    dataset_id = await _create_dataset(client, csv_upload)
    response = await _train(client, dataset_id)
    ranges = await client.get(f"/api/v1/runs/{response.json()['id']}/results/ranges")
    assert ranges.status_code == 404, ranges.text


def _two_target_training_csv() -> bytes:
    rows = "\n".join(
        f"{smiles},{1.0 + 0.37 * index},{(index // 2) % 2}"
        for index, smiles in enumerate(_STRUCTURES)
    )
    return f"smiles,y,active\n{rows}\n".encode()


async def test_a_two_target_protocol_writes_every_readout_and_its_own_uncertainty(
    client, csv_upload, prediction_upload_ref
):
    upload_ref = await csv_upload(_two_target_training_csv())
    dataset = await client.post(
        "/api/v1/datasets",
        json={
            "name": "panel",
            "upload_ref": upload_ref,
            "structure_column": "smiles",
            "targets": [
                {"column": "y", "kind": "numeric", "unit": "logS", "direction": "high"},
                {"column": "active", "kind": "binary"},
            ],
            "split": {"strategy": "random", "seed": 1},
        },
    )
    assert dataset.status_code == 201, dataset.text
    trained = await _train(client, dataset.json()["id"], engine_id="ecfp4-randomforest")
    assert trained.status_code == 202, trained.text
    protocol_id = (await client.get("/api/v1/protocols")).json()["items"][0]["id"]
    assert (await client.post(f"/api/v1/protocols/{protocol_id}/publish")).status_code == 204

    run_id = (await _predict(client, protocol_id, prediction_upload_ref)).json()["id"]
    response = await client.get(
        f"/api/v1/runs/{run_id}/results",
        params={"sort_by": "active_uncertainty", "sort_dir": "desc"},
    )
    assert response.status_code == 200, response.text
    row = response.json()["items"][0]
    assert set(row["readouts"]) == {"y", "active_probability", "active"}
    assert set(row["uncertainty"]) == {"y", "active"}
    # the random forest reports a spread for each target separately
    assert isinstance(row["uncertainty"]["y"], float)
    assert isinstance(row["uncertainty"]["active"], float)


async def test_a_tuned_cutoff_labels_the_predicted_classes(client, csv_upload):
    """The class column is the probability read at the cutoff training tuned, which is
    stored on the readout -- not a fixed 0.5."""
    dataset = await client.post(
        "/api/v1/datasets",
        json={
            "name": "reactivity",
            "upload_ref": await csv_upload(tunable_csv()),
            "structure_column": "smiles",
            "targets": [{"column": "y", "kind": "binary"}],
            "split": {"strategy": "random", "seed": 1},
        },
    )
    assert dataset.status_code == 201, dataset.text
    trained = await _train(client, dataset.json()["id"], tune_cutoffs=True)
    assert trained.status_code == 202, trained.text
    protocol_id = (await client.get("/api/v1/protocols")).json()["items"][0]["id"]
    assert (await client.post(f"/api/v1/protocols/{protocol_id}/publish")).status_code == 204

    protocol = (await client.get(f"/api/v1/protocols/{protocol_id}")).json()
    (readout,) = [r for r in protocol["readouts"] if r["type"] == "class"]
    assert readout["threshold"] is not None
    scorecard = (await client.get(f"/api/v1/protocols/{protocol_id}/scorecard")).json()
    assert scorecard[0]["cutoff"] == readout["threshold"]

    upload_ref = await csv_upload(
        b"smiles\nCCCCCCCCCCCCO\nCCCCCCCCCCCCN\nCCCCCCCCCCCCCCCCCCCCCCCCCCCCO\nCCCCCCCCCCCCCCCCN\n"
    )
    run_id = (await _predict(client, protocol_id, upload_ref)).json()["id"]
    items = (await client.get(f"/api/v1/runs/{run_id}/results")).json()["items"]
    assert len(items) == 4
    for item in items:
        probability = item["readouts"]["y_probability"]["value"]
        assert item["readouts"]["y"]["value"] == (
            1.0 if probability >= readout["threshold"] else 0.0
        )


async def test_the_stored_cutoff_not_half_decides_the_predicted_class(
    client, session_factory, csv_upload
):
    """A tuned cutoff happens to coincide with a probability the model emits, so the
    test above cannot tell it from 0.5. Here the stored threshold is set above every
    probability the model gives an amine: the class must follow it, not 0.5."""
    dataset = await client.post(
        "/api/v1/datasets",
        json={
            "name": "reactivity",
            "upload_ref": await csv_upload(tunable_csv()),
            "structure_column": "smiles",
            "targets": [{"column": "y", "kind": "binary"}],
            "split": {"strategy": "random", "seed": 1},
        },
    )
    assert (await _train(client, dataset.json()["id"])).status_code == 202
    protocol_id = (await client.get("/api/v1/protocols")).json()["items"][0]["id"]
    readouts = (await client.get(f"/api/v1/protocols/{protocol_id}")).json()["readouts"]
    assert [r["threshold"] for r in readouts] == [None, None]  # untuned: 0.5
    async with session_factory() as session:
        await session.execute(
            update(InSilicoProtocolModel)
            .where(InSilicoProtocolModel.id == uuid.UUID(protocol_id))
            .values(
                readouts=[
                    {**r, "threshold": 0.999999} if r["type"] == "class" else r for r in readouts
                ]
            )
        )
        await session.commit()
    assert (await client.post(f"/api/v1/protocols/{protocol_id}/publish")).status_code == 204

    upload_ref = await csv_upload(b"smiles\nCCCCCCCCCCCCN\nCCCCCCCCCCCCO\n")
    run_id = (await _predict(client, protocol_id, upload_ref)).json()["id"]
    items = (await client.get(f"/api/v1/runs/{run_id}/results")).json()["items"]
    amine = next(i for i in items if i["structure"] == "CCCCCCCCCCCCN")
    assert amine["readouts"]["y_probability"]["value"] > 0.9
    assert amine["readouts"]["y"]["value"] == 0.0


async def test_run_response_carries_the_protocol_id_for_a_prediction(
    client, published_protocol_id, prediction_upload_ref
):
    """Important 5 (Task 17 review): without this, a client holding only a
    run id has no way to construct `GET /protocols/{id}` to fetch the readout
    metadata above."""
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    assert submitted.json()["protocol_id"] == published_protocol_id

    polled = await client.get(f"/api/v1/runs/{submitted.json()['id']}")
    assert polled.json()["protocol_id"] == published_protocol_id


async def test_a_training_run_has_no_protocol_id_when_it_is_first_accepted(client, csv_upload):
    """The 202 is built from the Run as enqueued, and at that moment the
    Protocol genuinely does not exist yet. Null here is honest, not missing."""
    dataset_id = await _create_dataset(client, csv_upload)
    response = await _train(client, dataset_id)
    assert response.status_code == 202, response.text
    assert response.json()["protocol_id"] is None


async def test_a_training_run_carries_the_name_the_scientist_typed(
    client, published_protocol_id, prediction_upload_ref
):
    """The Protocols page lists a training run before its Protocol exists, and
    this name is all it has to call it by. A prediction run has none."""
    await _predict(client, published_protocol_id, prediction_upload_ref)

    trainings = (await client.get("/api/v1/runs?kind=training")).json()["items"]
    assert [item["name"] for item in trainings] == ["solubility model"]
    predictions = (await client.get("/api/v1/runs?kind=prediction")).json()["items"]
    assert [item["name"] for item in predictions] == [None]


async def test_polling_a_finished_training_run_yields_the_protocol_it_produced(client, csv_upload):
    """The transition the whole training screen depends on: submit, poll to
    `ready`, then follow `protocol_id` to the Scorecard. Before Run gained the
    column this was a dead end -- the id existed only inside a blob path."""
    dataset_id = await _create_dataset(client, csv_upload)
    run_id = (await _train(client, dataset_id)).json()["id"]

    polled = (await client.get(f"/api/v1/runs/{run_id}")).json()
    assert polled["status"] == "ready", polled
    assert polled["protocol_id"] is not None

    scorecard = await client.get(f"/api/v1/protocols/{polled['protocol_id']}/scorecard")
    assert scorecard.status_code == 200, scorecard.text


async def test_a_third_identical_request_after_two_failures_does_not_500(
    client, published_protocol_id, prediction_upload_ref
):
    """CRITICAL fix (Task 17 review): `find_by_cache_key` had no `.limit(1)`,
    so the moment two rows shared a cache_key, every subsequent lookup raised
    `MultipleResultsFound` -> a raw 500. Reachable by a user typo'ing
    `structure_column` and retrying the *identical* (still-typo'd) request --
    each FAILED attempt shares its cache_key with the last (a FAILED Run is
    never a cache hit, so a retry always creates a fresh row under the same
    key), so two retries is all it takes. Measured over HTTP by the reviewer:
    202, 202, then (before the fix) 500 on the third identical request.

    Deliberately reuses the *same* bad `structure_column` on every attempt --
    a request with a different one has a different cache_key (Task 17's own
    TDD caught that), so it would never collide and never reach the bug.
    """
    first = await _predict(
        client, published_protocol_id, prediction_upload_ref, structure_column="does-not-exist"
    )
    assert first.status_code == 202, first.text
    first_run = await client.get(f"/api/v1/runs/{first.json()['id']}")
    assert first_run.json()["status"] == "failed"

    second = await _predict(
        client, published_protocol_id, prediction_upload_ref, structure_column="does-not-exist"
    )
    assert second.status_code == 202, second.text
    second_run = await client.get(f"/api/v1/runs/{second.json()['id']}")
    assert second_run.json()["status"] == "failed"
    assert second.json()["id"] != first.json()["id"]  # a fresh row, same cache_key

    third = await _predict(
        client, published_protocol_id, prediction_upload_ref, structure_column="does-not-exist"
    )
    assert third.status_code == 202, third.text  # not 500


async def test_running_a_draft_protocol_is_a_409(
    client, trained_protocol_id, prediction_upload_ref
):
    response = await _predict(client, trained_protocol_id, prediction_upload_ref)
    assert response.status_code == 409, response.text


async def test_predicting_against_an_unknown_protocol_is_a_404(client, prediction_upload_ref):
    response = await _predict(
        client, "11111111-1111-1111-1111-111111111111", prediction_upload_ref
    )
    assert response.status_code == 404, response.text


async def test_predicting_with_an_unknown_upload_ref_is_a_404(client, published_protocol_id):
    response = await _predict(
        client, published_protocol_id, "11111111-1111-1111-1111-111111111111"
    )
    assert response.status_code == 404, response.text


async def test_a_missing_run_id_is_a_404(client):
    response = await client.get(f"/api/v1/runs/{uuid.uuid4()}")
    assert response.status_code == 404, response.text


async def test_workspace_id_in_the_body_is_rejected(
    client, published_protocol_id, prediction_upload_ref
):
    response = await _predict(
        client,
        published_protocol_id,
        prediction_upload_ref,
        workspace_id="00000000-0000-0000-0000-000000000001",
    )
    assert response.status_code == 422, response.text


async def test_viewer_cannot_submit_a_prediction(
    viewer_client, published_protocol_id, prediction_upload_ref
):
    response = await _predict(viewer_client, published_protocol_id, prediction_upload_ref)
    assert response.status_code == 403, response.text


async def test_viewer_can_poll_and_read_results(
    client, viewer_client, published_protocol_id, prediction_upload_ref
):
    """Reading is not gated behind editor -- only submitting and cancelling."""
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    run_id = submitted.json()["id"]

    assert (await viewer_client.get(f"/api/v1/runs/{run_id}")).status_code == 200
    assert (await viewer_client.get(f"/api/v1/runs/{run_id}/results")).status_code == 200


async def test_runs_are_scoped_to_the_callers_workspace(
    client, other_workspace_client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    run_id = submitted.json()["id"]

    assert (await other_workspace_client.get(f"/api/v1/runs/{run_id}")).status_code == 404
    assert (await other_workspace_client.get(f"/api/v1/runs/{run_id}/results")).status_code == 404
    assert (await other_workspace_client.post(f"/api/v1/runs/{run_id}/cancel")).status_code == 404
    assert (await other_workspace_client.get("/api/v1/runs")).json()["items"] == []


async def test_a_protocol_from_another_workspace_is_not_runnable(
    other_workspace_client, published_protocol_id
):
    """Design intent: a published Protocol is a shared asset within its own
    workspace, not across tenants -- resolved the same way every other
    Protocol read is (Task 16's `test_protocols_are_scoped_to_the_callers_workspace`).
    """
    upload = await other_workspace_client.post(
        "/api/v1/datasets/uploads", files={"file": ("data.csv", _PREDICTION_CSV, "text/csv")}
    )
    assert upload.status_code == 201, upload.text

    response = await _predict(
        other_workspace_client, published_protocol_id, upload.json()["upload_ref"]
    )
    assert response.status_code == 404, response.text


async def test_unauthenticated_request_is_rejected(anonymous_client):
    assert (await anonymous_client.post("/api/v1/runs", json={})).status_code == 401
    assert (await anonymous_client.get(f"/api/v1/runs/{uuid.uuid4()}")).status_code == 401


async def test_results_for_a_training_run_is_a_404(client, csv_upload):
    """`/results` is a prediction-shaped resource; a training Run's id has none."""
    dataset_id = await _create_dataset(client, csv_upload)
    response = await _train(client, dataset_id)
    assert response.status_code == 202, response.text

    results = await client.get(f"/api/v1/runs/{response.json()['id']}/results")
    assert results.status_code == 404, results.text


async def test_results_for_a_run_with_no_results_yet_is_a_409(
    client, session_factory, workspace_id
):
    run = Run(
        kind=RunKind.PREDICTION,
        workspace_id=workspace_id,
        requested_by=uuid.uuid4(),
        cache_key="k",
    )
    await SqlAlchemyRunRepository(session_factory).add(run)

    response = await client.get(f"/api/v1/runs/{run.id}/results")
    assert response.status_code == 409, response.text


async def test_cancel_flips_a_pending_run_to_cancelled(client, session_factory, workspace_id):
    run = Run(
        kind=RunKind.PREDICTION,
        workspace_id=workspace_id,
        requested_by=uuid.uuid4(),
        cache_key="k2",
    )
    await SqlAlchemyRunRepository(session_factory).add(run)

    response = await client.post(f"/api/v1/runs/{run.id}/cancel")
    assert response.status_code == 204, response.text
    assert response.content == b""

    polled = await client.get(f"/api/v1/runs/{run.id}")
    assert polled.json()["status"] == "cancelled"


async def test_cancelling_an_already_ready_run_is_a_409(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    run_id = submitted.json()["id"]

    response = await client.post(f"/api/v1/runs/{run_id}/cancel")
    assert response.status_code == 409, response.text


async def test_viewer_cannot_cancel(viewer_client, client, session_factory, workspace_id):
    run = Run(
        kind=RunKind.PREDICTION,
        workspace_id=workspace_id,
        requested_by=uuid.uuid4(),
        cache_key="k3",
    )
    await SqlAlchemyRunRepository(session_factory).add(run)

    response = await viewer_client.post(f"/api/v1/runs/{run.id}/cancel")
    assert response.status_code == 403, response.text


async def test_paging_through_more_results_than_the_limit_terminates(
    client, published_protocol_id, csv_upload
):
    csv = b"smiles\n" + b"\n".join(smiles.encode() for smiles in _STRUCTURES[:6]) + b"\n"
    upload_ref = await csv_upload(csv)
    submitted = await _predict(client, published_protocol_id, upload_ref)
    run_id = submitted.json()["id"]

    seen: list[str] = []
    cursor: str | None = None
    pages = 0
    while True:
        url = f"/api/v1/runs/{run_id}/results?limit=2"
        if cursor is not None:
            url += f"&cursor={cursor}"
        response = await client.get(url)
        assert response.status_code == 200, response.text
        page = response.json()
        pages += 1
        seen.extend(item["structure"] for item in page["items"])
        cursor = page["next_cursor"]
        if cursor is None:
            break
        assert pages < 10, f"pagination did not terminate: {pages} pages, saw {seen}"

    assert pages == 3
    assert len(seen) == 6
    assert len(set(seen)) == 6


async def test_listing_runs_returns_both_kinds_newest_first(
    client, published_protocol_id, prediction_upload_ref
):
    """Training a Protocol and then predicting with it leaves two Runs. Both
    are listed, and the prediction -- created second -- comes back first."""
    await _predict(client, published_protocol_id, prediction_upload_ref)

    items = (await client.get("/api/v1/runs")).json()["items"]
    assert [item["kind"] for item in items] == ["prediction", "training"]


async def test_listing_runs_filters_by_kind(client, published_protocol_id, prediction_upload_ref):
    """A prediction Run is what a user browses; a training Run belongs to its
    Protocol's history. Filtering server-side keeps a client from paging
    through the wrong kind to assemble one screen."""
    await _predict(client, published_protocol_id, prediction_upload_ref)

    predictions = (await client.get("/api/v1/runs?kind=prediction")).json()["items"]
    assert [item["kind"] for item in predictions] == ["prediction"]

    trainings = (await client.get("/api/v1/runs?kind=training")).json()["items"]
    assert [item["kind"] for item in trainings] == ["training"]


async def test_an_unknown_run_kind_is_a_422_not_a_silent_unfiltered_list(client):
    """Silently ignoring an unrecognised filter would hand back every Run as
    though the filter had matched them all."""
    response = await client.get("/api/v1/runs?kind=generation")
    assert response.status_code == 422, response.text


async def test_listing_runs_pages_with_an_opaque_cursor(
    client, published_protocol_id, prediction_upload_ref
):
    await _predict(client, published_protocol_id, prediction_upload_ref)

    first = (await client.get("/api/v1/runs?limit=1")).json()
    assert len(first["items"]) == 1
    assert first["next_cursor"] is not None

    second = (await client.get(f"/api/v1/runs?limit=1&cursor={first['next_cursor']}")).json()
    assert len(second["items"]) == 1
    assert second["items"][0]["id"] != first["items"][0]["id"]
    assert second["next_cursor"] is None


async def test_results_carry_their_row_id(client, published_protocol_id, prediction_upload_ref):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    response = await client.get(f"/api/v1/runs/{submitted.json()['id']}/results")
    assert response.status_code == 200, response.text
    assert [item["row_id"] for item in response.json()["items"]] == [0, 1, 2]


async def test_results_can_be_sorted_by_a_readout(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    run_id = submitted.json()["id"]
    protocol = (await client.get(f"/api/v1/protocols/{published_protocol_id}")).json()
    readout = protocol["readouts"][0]["name"]

    ascending = await client.get(f"/api/v1/runs/{run_id}/results?sort_by={readout}&sort_dir=asc")
    descending = await client.get(f"/api/v1/runs/{run_id}/results?sort_by={readout}&sort_dir=desc")
    assert ascending.status_code == 200, ascending.text
    assert descending.status_code == 200, descending.text

    def values(response):
        return [item["readouts"][readout]["value"] for item in response.json()["items"]]

    assert values(ascending) == sorted(values(ascending))
    assert values(descending) == list(reversed(values(ascending)))
    # Reordering must not renumber: the same compound keeps the same handle,
    # because that handle is what POST /collections stores.
    assert {item["row_id"] for item in descending.json()["items"]} == {0, 1, 2}


async def test_results_can_be_filtered_by_a_range(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    run_id = submitted.json()["id"]
    unfiltered = await client.get(f"/api/v1/runs/{run_id}/results")
    assert len(unfiltered.json()["items"]) == 3

    # No compound can be more than perfectly applicable, so this floor empties
    # the view -- which must be a well-formed empty page, not a 4xx.
    response = await client.get(
        f"/api/v1/runs/{run_id}/results?filters=" + quote('{"applicability": {"min": 1.1}}')
    )
    assert response.status_code == 200, response.text
    assert response.json()["items"] == []
    assert response.json()["next_cursor"] is None


async def test_an_unknown_filter_column_is_a_422(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    response = await client.get(
        f"/api/v1/runs/{submitted.json()['id']}/results?filters="
        + quote('{"not_a_column": {"min": 1}}')
    )
    assert response.status_code == 422, response.text


async def test_malformed_filter_json_is_a_422(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    response = await client.get(f"/api/v1/runs/{submitted.json()['id']}/results?filters=not-json")
    assert response.status_code == 422, response.text


async def test_a_filter_with_no_bounds_is_a_422(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    response = await client.get(
        f"/api/v1/runs/{submitted.json()['id']}/results?filters=" + quote('{"applicability": {}}')
    )
    assert response.status_code == 422, response.text


async def test_an_unknown_sort_direction_is_a_422(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(client, published_protocol_id, prediction_upload_ref)
    protocol = (await client.get(f"/api/v1/protocols/{published_protocol_id}")).json()
    response = await client.get(
        f"/api/v1/runs/{submitted.json()['id']}/results"
        f"?sort_by={protocol['readouts'][0]['name']}&sort_dir=sideways"
    )
    assert response.status_code == 422, response.text


async def test_an_identifier_column_is_carried_into_the_results(
    client, published_protocol_id, csv_upload
):
    """The join back to the scientist's own file: `compound_id` verbatim (blank
    becomes null), and `input_row` as the 1-based line in the upload."""
    upload_ref = await csv_upload(b"name,smiles\nCPD-1,CCO\n,CCC\nCPD-1,c1ccccc1\n")
    submitted = await _predict(client, published_protocol_id, upload_ref, id_column="name")
    assert submitted.status_code == 202, submitted.text
    run_id = submitted.json()["id"]
    # The POST body is the Run as created; inline jobs finish before the poll.
    run = (await client.get(f"/api/v1/runs/{run_id}")).json()
    assert run["status"] == "ready", run
    assert run["metrics"] == {"uploaded_rows": 3, "scored_rows": 3}

    page = (await client.get(f"/api/v1/runs/{run_id}/results")).json()
    assert [item["compound_id"] for item in page["items"]] == ["CPD-1", None, "CPD-1"]
    assert [item["input_row"] for item in page["items"]] == [1, 2, 3]


async def test_an_unknown_identifier_column_fails_the_run_by_name(
    client, published_protocol_id, prediction_upload_ref
):
    submitted = await _predict(
        client, published_protocol_id, prediction_upload_ref, id_column="nope"
    )
    assert submitted.status_code == 202, submitted.text
    run = (await client.get(f"/api/v1/runs/{submitted.json()['id']}")).json()
    assert run["status"] == "failed"
    assert "nope" in run["error_message"]


def _failed_prediction(workspace_id, protocol_id: str, upload_ref: str, cache_key: str) -> Run:
    run = Run(
        kind=RunKind.PREDICTION,
        workspace_id=workspace_id,
        requested_by=uuid.uuid4(),
        cache_key=cache_key,
        params={
            "protocol_id": protocol_id,
            "upload_ref": upload_ref,
            "structure_column": "smiles",
            "conditions": {},
        },
        protocol_id=uuid.UUID(protocol_id),
    )
    run.start()
    run.fail("the runner died")
    return run


async def test_retrying_a_failed_prediction_reenqueues_and_runs_it(
    client, session_factory, workspace_id, published_protocol_id, prediction_upload_ref
):
    run = _failed_prediction(workspace_id, published_protocol_id, prediction_upload_ref, "retry-1")
    await SqlAlchemyRunRepository(session_factory).add(run)

    response = await client.post(f"/api/v1/runs/{run.id}/retry")
    assert response.status_code == 204, response.text

    polled = (await client.get(f"/api/v1/runs/{run.id}")).json()
    # Inline jobs run inside the request, so the retried run has already finished --
    # on the same row, with its failure cleared.
    assert polled["status"] == "ready", polled
    assert polled["error_message"] is None


async def test_retrying_a_ready_run_is_a_409(client, published_protocol_id, prediction_upload_ref):
    run_id = (await _predict(client, published_protocol_id, prediction_upload_ref)).json()["id"]
    assert (await client.post(f"/api/v1/runs/{run_id}/retry")).status_code == 409


async def test_viewer_cannot_retry(
    viewer_client, session_factory, workspace_id, published_protocol_id, prediction_upload_ref
):
    run = _failed_prediction(workspace_id, published_protocol_id, prediction_upload_ref, "retry-2")
    await SqlAlchemyRunRepository(session_factory).add(run)
    assert (await viewer_client.post(f"/api/v1/runs/{run.id}/retry")).status_code == 403


async def _failed_training_with_saved_progress(
    app, session_factory, workspace_id, requested_by
) -> tuple[Run, list[str]]:
    """A failed training run with two blobs saved under its checkpoint root.

    The dataset it names does not exist, so a retried attempt fails at once, before it
    could save or clear anything itself: whatever is left under the root afterwards is
    what the retry did or did not discard.
    """
    dataset_id = uuid.uuid4()
    run = Run(
        kind=RunKind.TRAINING,
        workspace_id=workspace_id,
        requested_by=requested_by,
        cache_key="retry-training",
        params=TrainProtocolCommand(
            name="solubility model",
            dataset_id=dataset_id,
            engine_id="ecfp4-xgboost",
            conditions={},
        ).to_params(),
    )
    run.start()
    run.fail("the runner died")
    await SqlAlchemyRunRepository(session_factory).add(run)
    store = app.state.container[BlobStore]
    root = checkpoint_root(workspace_id, dataset_id, run.id)
    keys = [f"{root}model/fit.a", f"{root}model/fit.b"]
    for key in keys:
        store.put_bytes(key, b"saved progress")
    return run, keys


async def test_start_over_discards_the_saved_progress_before_requeueing(
    app, client, session_factory, workspace_id, client_user_id
):
    run, keys = await _failed_training_with_saved_progress(
        app, session_factory, workspace_id, client_user_id
    )

    response = await client.post(f"/api/v1/runs/{run.id}/retry", json={"fresh": True})
    assert response.status_code == 204, response.text

    store = app.state.container[BlobStore]
    assert [store.exists(key) for key in keys] == [False, False]
    # The requeued attempt ran (inline) and failed on the missing dataset, as designed.
    assert (await client.get(f"/api/v1/runs/{run.id}")).json()["status"] == "failed"


async def test_a_start_over_whose_delete_fails_leaves_the_run_stopped_and_its_progress(
    app, client, session_factory, workspace_id, client_user_id, monkeypatch
):
    run, keys = await _failed_training_with_saved_progress(
        app, session_factory, workspace_id, client_user_id
    )
    store = app.state.container[BlobStore]

    def unavailable(prefix: str) -> None:
        raise OSError("blob store unavailable")

    monkeypatch.setattr(store, "delete_prefix", unavailable)
    try:
        response = await client.post(f"/api/v1/runs/{run.id}/retry", json={"fresh": True})
        assert response.status_code >= 500
    except OSError:
        pass  # the test transport re-raises app exceptions instead of answering 500
    monkeypatch.undo()

    # Not requeued: resuming here would load exactly what the person asked to discard.
    stored = await SqlAlchemyRunRepository(session_factory).get_by_id(run.id)
    assert stored is not None and stored.status is RunStatus.FAILED
    assert [store.exists(key) for key in keys] == [True, True]


async def test_a_plain_retry_keeps_the_saved_progress(
    app, client, session_factory, workspace_id, client_user_id
):
    run, keys = await _failed_training_with_saved_progress(
        app, session_factory, workspace_id, client_user_id
    )

    response = await client.post(f"/api/v1/runs/{run.id}/retry")
    assert response.status_code == 204, response.text

    store = app.state.container[BlobStore]
    assert [store.exists(key) for key in keys] == [True, True]
    # Proof the retry did run the job: it failed again, on the missing dataset.
    assert (await client.get(f"/api/v1/runs/{run.id}")).json()["status"] == "failed"


async def test_start_over_on_a_prediction_run_changes_nothing_else(
    client, session_factory, workspace_id, published_protocol_id, prediction_upload_ref
):
    run = _failed_prediction(workspace_id, published_protocol_id, prediction_upload_ref, "retry-3")
    await SqlAlchemyRunRepository(session_factory).add(run)

    response = await client.post(f"/api/v1/runs/{run.id}/retry", json={"fresh": True})
    assert response.status_code == 204, response.text

    polled = (await client.get(f"/api/v1/runs/{run.id}")).json()
    assert polled["status"] == "ready", polled
    assert polled["error_message"] is None


async def test_a_retry_body_with_a_non_boolean_or_unknown_field_is_a_422(
    client, session_factory, workspace_id, published_protocol_id, prediction_upload_ref
):
    run = _failed_prediction(workspace_id, published_protocol_id, prediction_upload_ref, "retry-4")
    await SqlAlchemyRunRepository(session_factory).add(run)

    for body in ({"fresh": "yes"}, {"other": 1}):
        response = await client.post(f"/api/v1/runs/{run.id}/retry", json=body)
        assert response.status_code == 422, (body, response.text)
    # Refused before anything ran: the run is still failed.
    assert (await client.get(f"/api/v1/runs/{run.id}")).json()["status"] == "failed"


# The export: what the grid shows, as a workbook, with the upload's own columns.
_EXPORT_CSV = (
    b"smiles,name,measured,code,note\n"
    b'CCO,ethanol,1.5,007,=HYPERLINK("http://x")\n'
    b"not-a-smiles,broken,2.0,008,dropped\n"
    b"c1ccccc1,benzene,-0.25,009,plain\n"
    b"Fc1ccc(F)cc1,difluoro,NA,010,\n"
)


def _sheet(content: bytes, name: str) -> list[list[object]]:
    from io import BytesIO

    from openpyxl import load_workbook

    book = load_workbook(BytesIO(content))
    return [[cell.value for cell in row] for row in book[name].iter_rows()]


async def test_the_export_holds_the_grids_rows_and_the_uploads_own_columns(
    client, published_protocol_id, csv_upload
):
    upload_ref = await csv_upload(_EXPORT_CSV)
    run = await _predict(client, published_protocol_id, upload_ref, id_column="name")
    run_id = run.json()["id"]
    params = {"sort_by": "y", "sort_dir": "desc"}
    response = await client.get(f"/api/v1/runs/{run_id}/results/export", params=params)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    header, *rows = _sheet(response.content, "Predictions")
    assert header == [
        "ID",
        "Row",
        "SMILES",
        "y (logS, higher is better)",
        "Uncertainty",
        "Applicability",
        "measured",
        "code",
        "note",
    ]
    # The grid's own rows, in the grid's own order: the unparseable one is not scored.
    grid = (await client.get(f"/api/v1/runs/{run_id}/results", params=params)).json()["items"]
    assert [row[0] for row in rows] == [item["compound_id"] for item in grid]
    assert [row[1] for row in rows] == [item["input_row"] for item in grid]
    by_id = {row[0]: row for row in rows}
    # Matched back by row: numbers stay numbers, a leading-zero code stays text,
    # "NA" leaves the column as text, and a formula stays the text it was.
    assert by_id["benzene"][6] == "-0.25"
    assert by_id["benzene"][7] == "009"
    assert by_id["ethanol"][8] == '=HYPERLINK("http://x")'

    about = _sheet(response.content, "About")
    facts = {row[0]: row[1] for row in about if row and row[0]}
    assert facts["Compounds in this file"] == "3 of 3 scored"
    assert facts["Sorted by"] == "y (logS, higher is better), highest first"
    assert facts["Filters"] == "None"


async def test_the_export_keeps_numeric_upload_columns_numeric_and_applies_filters(
    client, published_protocol_id, csv_upload
):
    upload_ref = await csv_upload(b"smiles,measured\nCCO,1.5\nc1ccccc1,-0.25\nFc1ccc(F)cc1,3\n")
    run_id = (await _predict(client, published_protocol_id, upload_ref)).json()["id"]
    filters = '{"applicability": {"min": 0.0, "max": 1.0}}'
    response = await client.get(
        f"/api/v1/runs/{run_id}/results/export", params={"filters": filters}
    )
    assert response.status_code == 200, response.text
    header, *rows = _sheet(response.content, "Predictions")
    assert "ID" not in header
    assert [row[header.index("measured")] for row in rows] == [1.5, -0.25, 3.0]
    facts = {row[0]: row[1] for row in _sheet(response.content, "About") if row and row[0]}
    assert facts["Filters"] == "Applicability: 0 to 1"


async def test_a_formula_in_an_upload_is_written_as_text_not_a_formula(
    client, published_protocol_id, csv_upload
):
    from io import BytesIO
    from zipfile import ZipFile

    upload_ref = await csv_upload(b"smiles,note\nCCO,=1+1\n")
    run_id = (await _predict(client, published_protocol_id, upload_ref)).json()["id"]
    response = await client.get(f"/api/v1/runs/{run_id}/results/export")
    sheet_xml = ZipFile(BytesIO(response.content)).read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in sheet_xml


async def test_exporting_a_training_run_is_a_404(client, csv_upload):
    dataset_id = await _create_dataset(client, csv_upload)
    response = await _train(client, dataset_id)
    export = await client.get(f"/api/v1/runs/{response.json()['id']}/results/export")
    assert export.status_code == 404, export.text
