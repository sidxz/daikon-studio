"""A prediction Run's results as an Excel workbook.

The Predictions sheet holds what the triage grid showed, under the same filters and
sort, with the scientist's own upload columns beside it (matched on `input_row`). The
About sheet says where the numbers came from and what each column means, in plain
words, because a workbook travels to people who never saw the grid.
"""

from __future__ import annotations

import asyncio
import io
import uuid
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import UTC
from typing import Any

import polars as pl
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from returns.pipeline import is_successful
from returns.result import Failure, Result, Success

from daikonstudio.application.auth import AuthContext
from daikonstudio.application.catalog.derive_readouts import target_columns_of
from daikonstudio.application.data.create_dataset import upload_key
from daikonstudio.application.data.export_collection import _unit_and_direction
from daikonstudio.application.data.prepare_frame import read_csv_upload
from daikonstudio.application.execution.build_scorecard import _APPLICABILITY_THRESHOLD
from daikonstudio.application.execution.predict_with_protocol import (
    PredictWithProtocolCommand,
    load_results,
    view_results,
)
from daikonstudio.application.execution.result_view import RangeFilter, SortSpec
from daikonstudio.application.ports.blob_store import BlobStore
from daikonstudio.application.ports.protocol_repository import ProtocolRepository
from daikonstudio.application.ports.run_repository import RunRepository
from daikonstudio.domain.catalog.protocol import InSilicoProtocol
from daikonstudio.domain.catalog.readout import Readout, ReadoutType
from daikonstudio.domain.data.target import uncertainty_column
from daikonstudio.domain.execution.run import Run
from daikonstudio.domain.shared.errors import DomainError, ValidationError

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: A worksheet's last row, less the header.
_EXCEL_MAX_ROWS = 1_048_575
#: Control characters are not legal in the XML a worksheet is written as.
_ILLEGAL = r"[\x00-\x08\x0b\x0c\x0e-\x1f]"
_JOIN_KEY = "__export_row__"


@dataclass(frozen=True, kw_only=True)
class ExportPredictionResultsQuery:
    run_id: uuid.UUID
    sort: SortSpec | None = None
    filters: tuple[RangeFilter, ...] = ()


class ExportPredictionResults:
    def __init__(
        self, runs: RunRepository, protocols: ProtocolRepository, store: BlobStore
    ) -> None:
        self._runs = runs
        self._protocols = protocols
        self._store = store

    async def __call__(
        self, query: ExportPredictionResultsQuery, auth: AuthContext | None = None
    ) -> Result[bytes, DomainError]:
        loaded = await load_results(self._runs, self._protocols, self._store, query.run_id, auth)
        if not is_successful(loaded):
            return Failure(loaded.failure())
        run, protocol, frame = loaded.unwrap()

        viewed = view_results(protocol, frame, sort=query.sort, filters=query.filters)
        if not is_successful(viewed):
            return Failure(viewed.failure())
        rows = viewed.unwrap()
        if rows.height > _EXCEL_MAX_ROWS:
            return Failure(
                ValidationError(
                    f"These results have {rows.height:,} rows; an Excel sheet holds at most "
                    f"{_EXCEL_MAX_ROWS:,}.",
                    detail="Filter the results down before exporting.",
                )
            )

        upload, upload_note = self._upload(run, rows)
        # Thousands of rows through openpyxl is seconds of CPU: off the event loop,
        # as `GetScorecard` does for its similarity search.
        content = await asyncio.to_thread(
            _workbook,
            run=run,
            protocol=protocol,
            rows=rows,
            total=frame.height,
            upload=upload,
            upload_note=upload_note,
            sort=query.sort,
            filters=query.filters,
        )
        return Success(content)

    def _upload(self, run: Run, rows: pl.DataFrame) -> tuple[pl.DataFrame | None, str | None]:
        """The upload's own columns, keyed by row, or None and why not."""
        if "input_row" not in rows.columns:
            return None, "Not included: this run predates row numbers, so rows cannot be matched."
        command = PredictWithProtocolCommand.from_params(run.params)
        key = upload_key(run.workspace_id, uuid.UUID(command.upload_ref))
        try:
            upload = read_csv_upload(self._store.get_bytes(key))
        except FileNotFoundError:
            return None, "Not included: the uploaded file is no longer stored."
        except ValidationError:
            return None, "Not included: the uploaded file could not be read again."
        # Already in the sheet, as SMILES (as Studio read it) and ID.
        dropped = {command.structure_column, command.id_column}
        upload = upload.select([c for c in upload.columns if c not in dropped])
        if not upload.columns:
            return None, None
        return upload.select([_typed(upload[c]) for c in upload.columns]).with_row_index(
            _JOIN_KEY, offset=1
        ), None


def _typed(column: pl.Series) -> pl.Series:
    """A column of plain numbers as numbers, so it sorts and sums in Excel; anything
    else, and any number written with a leading zero (an ID like "00123"), as typed."""
    numbers = column.cast(pl.Float64, strict=False)
    text = column.drop_nulls()
    if (
        numbers.null_count() == column.null_count()
        # "inf" and "nan" parse, but a worksheet cannot hold either.
        and numbers.drop_nulls().is_finite().all()
        and not text.str.contains(r"^[+-]?0\d").any()
    ):
        return numbers
    return column.str.replace_all(_ILLEGAL, "")


def _labels(protocol: InSilicoProtocol, rows: pl.DataFrame) -> dict[str, str]:
    """Results column -> header, in the grid's order; what the sheet holds."""
    labels: dict[str, str] = {}
    if "compound_id" in rows.columns and rows["compound_id"].drop_nulls().len() > 0:
        labels["compound_id"] = "ID"
    if "input_row" in rows.columns:
        labels["input_row"] = "Row"
    labels["structure"] = "SMILES"
    for readout in protocol.readouts:
        suffix = _unit_and_direction(readout)
        labels[readout.name] = f"{readout.name} ({suffix})" if suffix else readout.name
    targets = target_columns_of(protocol.readouts)
    for target in targets:
        name = "Uncertainty" if len(targets) == 1 else f"Uncertainty ({target})"
        labels[uncertainty_column(target, target_count=len(targets))] = name
    labels["applicability"] = "Applicability"
    return labels


def _meanings(
    protocol: InSilicoProtocol, labels: dict[str, str], *, from_chemcellar: bool = False
) -> list[tuple[str, str]]:
    readouts = {readout.name: readout for readout in protocol.readouts}
    meanings: list[tuple[str, str]] = []
    for column, label in labels.items():
        readout = readouts.get(column)
        if column == "compound_id":
            text = (
                "The compound's ChemCellar registration number."
                if from_chemcellar
                else "Your identifier for the compound, from the uploaded file."
            )
        elif column == "input_row":
            text = (
                "The compound's position in the list imported from ChemCellar."
                if from_chemcellar
                else "The compound's row in your uploaded file, not counting the header."
            )
        elif column == "structure":
            text = "The compound's structure, as Studio read it."
        elif column == "applicability":
            text = (
                "How similar the compound is to the ones the model learned from, from 0 to 1. "
                f"Below {_APPLICABILITY_THRESHOLD}, the model has seen nothing like it, so treat "
                "the prediction as a rough guess."
            )
        elif readout is None:
            text = (
                "How unsure the model is about this compound; higher means less sure. "
                "Empty when the model does not report it."
            )
        elif readout.type is ReadoutType.PROBABILITY:
            text = "The model's estimated chance that the compound is active, from 0 to 1."
        elif readout.type is ReadoutType.CLASS:
            cutoff = _cutoff(readout.threshold if readout.threshold is not None else 0.5)
            text = f"Active when the estimated chance is at least {cutoff}, otherwise Inactive."
        else:
            unit = f", in {readout.unit}" if readout.unit else ""
            text = f"The model's predicted {readout.name}{unit}."
        meanings.append((label, text))
    return meanings


def _source_line(params: Mapping[str, Any]) -> str | None:
    """Where the compounds came from, when not from an uploaded file."""
    source = params.get("source")
    if not source:
        return None
    return f"ChemCellar: {source['protocol_name']}, run of {source['run_date']}"


def _cutoff(value: float) -> str:
    """Three significant digits, as the scorecard shows a cutoff (`formatCutoff`), but
    never rounded up to 1: "at least 1" for 0.99997 would mean nothing is ever active."""
    text = f"{value:.3g}"
    if float(text) < 1 or value >= 1:
        return text
    for digits in range(3, 9):
        text = f"{value:.{digits}f}"
        if float(text) < 1:
            return text
    return str(value)


def _bounds(range_filter: RangeFilter) -> str:
    low, high = range_filter.minimum, range_filter.maximum
    if low is not None and high is not None:
        return f"{low:g} to {high:g}"
    return f"at least {low:g}" if low is not None else f"at most {high:g}"


def _workbook(
    *,
    run: Run,
    protocol: InSilicoProtocol,
    rows: pl.DataFrame,
    total: int,
    upload: pl.DataFrame | None,
    upload_note: str | None,
    sort: SortSpec | None,
    filters: tuple[RangeFilter, ...],
) -> bytes:
    labels = _labels(protocol, rows)
    readouts = {readout.name: readout for readout in protocol.readouts}
    sheet = rows.select(
        [_shown(column, rows.schema[column], readouts.get(column)) for column in labels]
    ).rename(labels)

    uploaded: list[str] = []
    if upload is not None:
        # An upload column sharing a header with a results column keeps its data under
        # its own marked name, rather than one silently replacing the other.
        renamed = {
            c: f"{c} (from upload)" if c in sheet.columns else c
            for c in upload.columns
            if c != _JOIN_KEY
        }
        uploaded = list(renamed.values())
        sheet = (
            sheet.with_columns(rows["input_row"].cast(pl.Int64).alias(_JOIN_KEY))
            .join(
                upload.rename(renamed).with_columns(pl.col(_JOIN_KEY).cast(pl.Int64)),
                on=_JOIN_KEY,
                how="left",
                maintain_order="left",
            )
            .drop(_JOIN_KEY)
        )

    book = Workbook(write_only=True)
    predictions = book.create_sheet("Predictions")
    predictions.freeze_panes = "A2"
    for index, header in enumerate(sheet.columns, start=1):
        width = 48 if header == "SMILES" else max(10, min(40, len(header) + 2))
        predictions.column_dimensions[get_column_letter(index)].width = width
    predictions.append([_bold(predictions, header) for header in sheet.columns])
    for values in sheet.iter_rows():
        predictions.append([_cell(predictions, value) for value in values])

    about = book.create_sheet("About")
    about.column_dimensions["A"].width = 36
    about.column_dimensions["B"].width = 100
    sorted_by = (
        "Upload order"
        if sort is None
        else f"{labels[sort.column]}, {'highest' if sort.descending else 'lowest'} first"
    )
    source = _source_line(run.params)
    created = f"{run.created_at.astimezone(UTC):%Y-%m-%d %H:%M} UTC" if run.created_at else ""
    for line in [
        [_bold(about, "Predictions from DAIKON Studio")],
        [],
        ["Protocol", protocol.name],
        ["Engine", protocol.engine_id],
        ["Run started", created],
        *([["Compounds from", source]] if source else []),
        ["Compounds in this file", f"{rows.height:,} of {total:,} scored"],
        ["Filters", "; ".join(f"{labels[f.column]}: {_bounds(f)}" for f in filters) or "None"],
        ["Sorted by", sorted_by],
        *([["Your uploaded columns", upload_note]] if upload_note else []),
        [],
        [_bold(about, "Column"), _bold(about, "What it means")],
        *[
            [_cell(about, label), text]
            for label, text in _meanings(protocol, labels, from_chemcellar=bool(source))
        ],
        *[
            [_cell(about, name), "From ChemCellar." if source else "From your uploaded file."]
            for name in uploaded
        ],
    ]:
        about.append(line)

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _shown(column: str, dtype: pl.DataType, readout: Readout | None) -> pl.Expr:
    """A results column as the sheet shows it: a class as the grid's word, NaN as an
    empty cell, text without the control characters a worksheet cannot hold."""
    value = pl.col(column)
    if readout is not None and readout.type is ReadoutType.CLASS:
        # Already 1 or 0, decided at the Protocol's own cutoff when the run was scored
        # (`RunPrediction`); this only names it. Null stays empty.
        active, inactive = pl.lit("Active"), pl.lit("Inactive")
        return pl.when(value == 1.0).then(active).when(value == 0.0).then(inactive).alias(column)
    if dtype.is_float():
        return value.fill_nan(None)
    if dtype == pl.String:
        return value.str.replace_all(_ILLEGAL, "")
    return value


def _bold(sheet: Any, value: str) -> Any:
    cell = WriteOnlyCell(sheet, value=value)
    if value.startswith("="):
        cell.data_type = "s"
    cell.font = Font(bold=True)
    return cell


def _cell(sheet: Any, value: object) -> Any:
    """openpyxl turns any string starting with "=" into a live formula. An uploaded
    value must stay the text it was, never a formula that runs when the file opens."""
    if isinstance(value, str) and value.startswith("="):
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        return cell
    return value
